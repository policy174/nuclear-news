"""원자력 질의 구간 → 회사 국감 모니터링 양식(주제·의원 / 주요 질의 / 주요 답변 / 스크립트).

자막엔 발언자 이름이 없다. 위원장 호명("다음은 ○○○ 위원")과 문맥으로 Gemini 가 추정하고,
확실하지 않으면 '의원'·'장관' 같은 직함만 쓴다. 정당은 자막에 나온 경우만.

구간 = 원자력 알림 줄들을 10분 간격으로 묶은 것. 마지막 알림 뒤 5분이 지나면 '닫힘'으로 보고
한 번만 요약해 sections.json 에 캐시한다(같은 구간을 30초마다 다시 부르지 않게).
"""
import json
import os
import time
from datetime import datetime

import requests

GAP_SEC = 600          # 알림 사이가 이보다 벌어지면 다른 구간
CLOSE_SEC = 300        # 마지막 알림 뒤 이만큼 자막이 더 쌓이면 구간 닫힘
PAD_BEFORE_SEC = 240   # 구간 앞뒤로 붙일 자막 (호명·답변 포함)
PAD_AFTER_SEC = 300
MODELS = os.environ.get("AUDIT_GEMINI_MODELS", "gemini-3.5-flash,gemini-flash-latest,gemini-2.5-flash,gemini-3.1-flash-lite").split(",")
RETRY_AFTER_SEC = 120  # 전부 실패하면 이만큼 쉬고 재시도 (30초 렌더마다 두드리지 않게)
_next_try = [0.0]

PROMPT = """한국수력원자력 원자력정책실의 국정감사 모니터링 보고 초안을 만든다.
입력은 국회 의사중계 AI 자막(음성인식, 오탈자 있음)의 한 구간이다. [시각] 문장 형식.

출력 JSON:
{"title": "질의 주제 한 줄(예: 신규원전 백지화 및 원전 생태계 위축)",
 "member": "질의 의원 이름(위원장 호명·자기소개로 확인될 때만, 아니면 빈 문자열)",
 "party": "정당(자막에 나올 때만, 아니면 빈 문자열)",
 "answerer": "답변자 '이름 직함'(예: 김성환 장관). 이름을 모르면 직함만(예: 장관, 사장)",
 "questions": ["주요 질의 요지 개조식 1~4개"],
 "answers": ["주요 답변 요지 개조식 1~3개, 답변이 없으면 빈 배열"],
 "script": [{"speaker": "이현성 의원 | 김성환 장관 | 위원장 | 의원 | 장관", "text": "발언"}]}

규칙
- 요지는 개조식 체언 종결('~ 요구', '~ 지적', '~ 입장'). '~를 위한/~에 따른'으로 수식 관계를 살려 명사 무더기 금지.
- 수치·연도·호기·주체를 자막 그대로 정확히('16년, 5.5조원, 고리1호기). 자막에 없는 내용 추가 금지.
- 질의에는 취지(왜 묻는지)를 한 구절 담는다.
- script 는 원자력·한수원·전력정책에 해당하는 문답만, 말한 순서대로. 음성인식 오탈자는 문맥상 확실할 때만
  고치고 말투는 살린다. 같은 사람 말이 끊겼으면 이어 붙인다.
- 발언자를 모르면 지어내지 말고 '의원'·'장관'으로.
- 원자력과 무관한 구간이면 {"skip": true}."""


def blocks(alerts, last_line_ts, final=False):
    """alerts: line_time(ISO) 정렬 목록 → [(start_ts, end_ts, closed)]"""
    out = []
    times = sorted(datetime.fromisoformat(a["line_time"]).timestamp() for a in alerts)
    for t in times:
        if out and t - out[-1][1] <= GAP_SEC:
            out[-1][1] = t
        else:
            out.append([t, t])
    return [(s, e, final or last_line_ts - e >= CLOSE_SEC) for s, e in out]


def window(lines, start, end):
    lo, hi = start - PAD_BEFORE_SEC, end + PAD_AFTER_SEC
    return [(ts, tx) for ts, tx in lines if lo <= datetime.fromisoformat(ts).timestamp() <= hi]


def summarize(win):
    key = os.environ.get("GEMINI_API_KEY")
    if not key:
        return None
    text = "\n".join(f"[{ts[11:19]}] {tx}" for ts, tx in win)[:30000]
    body = {"systemInstruction": {"parts": [{"text": PROMPT}]},
            "contents": [{"role": "user", "parts": [{"text": text}]}],
            "generationConfig": {"responseMimeType": "application/json", "temperature": 0.2,
                                 "thinkingConfig": {"thinkingBudget": 0}}}
    if time.time() < _next_try[0]:
        return None
    for model in MODELS:   # 무료 키는 모델별 한도 — 429 면 다음 모델
        try:
            r = requests.post(f"https://generativelanguage.googleapis.com/v1beta/models/{model.strip()}:generateContent",
                              params={"key": key}, json=body, timeout=120)
            r.raise_for_status()
            out = json.loads(r.json()["candidates"][0]["content"]["parts"][0]["text"])
            if isinstance(out, dict):
                return out
        except Exception as e:   # 키 값이 섞일 수 있는 URL 은 찍지 않는다
            print("요약 실패", model, type(e).__name__, str(e).split(" for url")[0][:120])
    _next_try[0] = time.time() + RETRY_AFTER_SEC
    return None


def update_sections(cache_path, lines, alerts, final=False, summarizer=summarize):
    """닫힌 구간 중 아직 요약 안 된 것만 요약. → 구간 목록(열린 구간은 status=open, 요약 없음)"""
    cache = json.loads(cache_path.read_text(encoding="utf-8")) if cache_path.exists() else {}
    if not lines:
        return []
    last = datetime.fromisoformat(lines[-1][0]).timestamp()
    out = []
    for s, e, closed in blocks(alerts, last, final):
        k = f"{int(s)}-{int(e)}"
        win = window(lines, s, e)
        if closed and k not in cache:
            res = summarizer(win)
            if res is None:   # 실패는 캐시 안 함 → 다음 회차 재시도, 그동안은 원문만
                out.append({"status": "failed", "start": win[0][0] if win else "", "lines": win})
                continue
            cache[k] = res
            cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
        if k in cache:
            if not cache[k].get("skip"):
                out.append({"status": "done", "start": win[0][0] if win else "", **cache[k]})
        else:
            out.append({"status": "open", "start": win[0][0] if win else "", "lines": win})
    return out


def add_sections(d, sections, date_str):
    """docx 에 회사 양식으로 구간을 쓴다. d = python-docx Document."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt, RGBColor, Cm

    BLUE = RGBColor(0x1F, 0x3A, 0x8A)
    y, m, dd = date_str[:10].split("-") if len(date_str) >= 10 else ("", "", "")
    wd = "월화수목금토일"[datetime.fromisoformat(date_str[:10]).weekday()] if dd else ""
    date_label = f"'{y[2:]}.{int(m)}.{int(dd)}.({wd})" if dd else ""

    def bullets(items):
        for it in items:
            p = d.add_paragraph()
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.6)
            p.add_run("□ " + it)

    def script_line(speaker, text):
        p = d.add_paragraph()
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(3.2), Cm(-3.2)
        r = p.add_run(f"({speaker}) ")
        if "의원" not in speaker and "위원장" not in speaker:
            r.font.color.rgb = BLUE
        p.add_run(text)

    for i, s in enumerate(sections):
        if i:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        who = " ".join(x for x in (s.get("party", ""), s.get("member", "") and s["member"] + " 의원") if x)
        title = s.get("title") or ("(요약 대기 — 질의 진행 중)" if s["status"] == "open" else "(요약 실패 — 원문)")
        h = d.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hr = h.add_run(f"{title} ({who})" if who else title)
        hr.bold, hr.font.size = True, Pt(14)
        dp = d.add_paragraph()
        dp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        dp.add_run(f"{date_label} {s.get('start', '')[11:16]}~")
        if s["status"] == "done":
            d.add_paragraph().add_run(f"[주요 질의 내용 - {s.get('member') and s['member'] + ' 의원' or '의원'}]").bold = True
            bullets(s.get("questions") or [])
            if s.get("answers"):
                d.add_paragraph().add_run(f"[주요 답변 내용 - {s.get('answerer') or '답변자'}]").bold = True
                bullets(s["answers"])
            d.add_paragraph()
            d.add_paragraph().add_run("[스크립트]").bold = True
            merged = []   # 같은 사람 말이 줄마다 끊겨 오면 한 덩어리로
            for line in s.get("script") or []:
                if not isinstance(line, dict):
                    continue
                sp, tx = line.get("speaker", "발언자"), line.get("text", "")
                if merged and merged[-1][0] == sp:
                    merged[-1][1] += " " + tx
                else:
                    merged.append([sp, tx])
            for sp, tx in merged:
                script_line(sp, tx)
        else:
            d.add_paragraph().add_run("[자막 원문]").bold = True
            for ts, tx in s.get("lines", []):
                script_line(ts[11:19], tx)
