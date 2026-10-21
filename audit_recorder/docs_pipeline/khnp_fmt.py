"""회사 국감 모니터링 양식(10-07 지니 제공 사진) 재현.

제목(가운데 굵게) / 날짜·시각(오른쪽) / [주제](굵게) / (HH:MM)(굵게) /
(이 종 배[국]) 개조식 발언 — 정당 배지는 색 바탕 흰 글자 한 글자, 이름은 글자 사이 띄움, 내어쓰기.
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

FONT = "한수원 한울림"
PARTY = {"국민의힘": ("국", "E61E2B"), "더불어민주당": ("민", "1E4FA3"), "조국혁신당": ("조", "0A2D6E"), "진보당": ("진", "D6001C")}
LABEL_CM = 3.4   # 화자 라벨 폭(내어쓰기)


def new_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.6
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(1.8)
    return d


def _run(p, text, bold=False, size=None, color=None):
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _badge(p, party):
    ch, fill = PARTY.get(party, ("?", "777777"))
    r = _run(p, ch, bold=True, size=9, color="FFFFFF")
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
    r._element.get_or_add_rPr().append(shd)


def title(d, text, when):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, text, bold=True, size=12.5)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(p, when)


def topic(d, text):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(14)
    _run(p, f"[{text}]", bold=True)


def stamp(d, hhmm):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(6)
    _run(p, f"({hhmm})", bold=True)


def say(d, who, text, party=None):
    """who: '이종배' + party → (이 종 배[국]) / 직함이면 그대로 (기후부장관)."""
    p = d.add_paragraph()
    p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(LABEL_CM), Cm(-LABEL_CM)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(LABEL_CM))
    if party:
        _run(p, "(" + " ".join(who) + " ")
        _badge(p, party)
        _run(p, ")\t")
    else:
        _run(p, f"({who})\t")
    _run(p, text)


def page(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
